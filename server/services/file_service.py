"""File ingestion: safe upload, type detection, text extraction, ZIP handling."""

import os
import re
import uuid
import zipfile
import logging
from typing import Any, BinaryIO
from datetime import datetime, timezone

from server.storage.db import query as db_query, execute as db_execute

log = logging.getLogger(__name__)

_SERVER_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_UPLOAD_DIR = os.environ.get("UPLOAD_DIR", "").strip()
if _UPLOAD_DIR:
    UPLOAD_DIR = os.path.abspath(_UPLOAD_DIR)
else:
    UPLOAD_DIR = os.path.join(_SERVER_ROOT, "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

MAX_FILE_SIZE = 50 * 1024 * 1024          # 50 MB per file
MAX_ZIP_ENTRY_SIZE = 5 * 1024 * 1024      # 5 MB per file inside a ZIP
MAX_ZIP_ENTRIES = 3000
MAX_ZIP_TOTAL_UNCOMPRESSED = 300 * 1024 * 1024  # zip-bomb guard
MAX_TEXT_EXTRACT_CHARS = 400_000

# Files we refuse to ingest unless the user explicitly allows them.
SENSITIVE_NAME_PATTERNS = (
    ".env",
    "id_rsa",
    "id_dsa",
    "id_ecdsa",
    "id_ed25519",
    "credentials",
    "secrets",
    "secret.",
    ".npmrc",
    ".pypirc",
    ".htpasswd",
    "service-account",
    "serviceaccount",
    "private-key",
    "privatekey",
)
SENSITIVE_EXTENSIONS = {".pem", ".key", ".p12", ".pfx", ".keystore", ".jks", ".asc", ".ppk"}

# Directories/files that are never useful as AI context.
SKIP_DIR_PARTS = (
    "node_modules", "__pycache__", ".git", ".hg", ".svn", ".venv", "venv",
    "site-packages", "dist-info", "egg-info", ".next", ".nuxt", ".turbo",
    ".pytest_cache", ".mypy_cache", ".ruff_cache", "coverage", ".idea", ".vscode",
    "target", "build", "dist", ".gradle", ".terraform", ".serverless",
)
SKIP_EXTENSIONS = {
    ".pyc", ".pyo", ".class", ".o", ".obj", ".so", ".dylib", ".dll", ".exe",
    ".bin", ".dat", ".lock", ".map", ".min.js", ".min.css", ".woff", ".woff2",
    ".ttf", ".eot", ".otf", ".mp3", ".mp4", ".mov", ".avi", ".mkv", ".wav",
    ".psd", ".ai", ".sketch", ".db", ".sqlite", ".sqlite3", ".pdb", ".iso",
}

CODE_EXTENSIONS = {
    ".js", ".jsx", ".mjs", ".cjs", ".ts", ".tsx", ".py", ".pyi", ".java",
    ".c", ".h", ".cpp", ".cc", ".cxx", ".hpp", ".hh", ".cs", ".go", ".rs",
    ".php", ".rb", ".sql", ".html", ".htm", ".css", ".scss", ".sass", ".less",
    ".sh", ".bash", ".zsh", ".ps1", ".psm1", ".bat", ".json", ".jsonc",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".xml", ".gradle",
    ".vue", ".svelte", ".astro", ".kt", ".kts", ".swift", ".dart", ".lua",
    ".r", ".scala", ".clj", ".cljs", ".ex", ".exs", ".erl", ".hrl", ".elm",
    ".hs", ".ml", ".mli", ".fs", ".fsx", ".pl", ".pm", ".groovy", ".tf",
    ".tfvars", ".proto", ".graphql", ".gql", ".prisma", ".ipynb", ".dockerfile",
}

TEXT_EXTENSIONS = {".txt", ".md", ".mdx", ".rst", ".csv", ".tsv", ".log", ".env.example"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".tif", ".svg"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".xls", ".odt", ".rtf"}
ARCHIVE_EXTENSIONS = {".zip"}

_MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".json": "application/json",
    ".csv": "text/csv",
    ".yaml": "text/yaml",
    ".yml": "text/yaml",
    ".xml": "application/xml",
    ".py": "text/x-python",
    ".js": "text/javascript",
    ".ts": "text/typescript",
    ".tsx": "text/typescript",
    ".html": "text/html",
    ".css": "text/css",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
    ".svg": "image/svg+xml",
    ".zip": "application/zip",
}

LANGUAGE_BY_EXT = {
    ".py": "python", ".pyi": "python", ".ipynb": "python",
    ".js": "javascript", ".mjs": "javascript", ".cjs": "javascript", ".jsx": "jsx",
    ".ts": "typescript", ".tsx": "tsx",
    ".java": "java", ".kt": "kotlin", ".kts": "kotlin",
    ".c": "c", ".h": "c", ".cpp": "cpp", ".cc": "cpp", ".cxx": "cpp",
    ".hpp": "cpp", ".hh": "cpp", ".cs": "csharp",
    ".go": "go", ".rs": "rust", ".php": "php", ".rb": "ruby",
    ".sql": "sql", ".html": "html", ".htm": "html",
    ".css": "css", ".scss": "scss", ".sass": "sass", ".less": "less",
    ".sh": "bash", ".bash": "bash", ".zsh": "bash", ".ps1": "powershell",
    ".bat": "batch", ".json": "json", ".jsonc": "json",
    ".yaml": "yaml", ".yml": "yaml", ".toml": "toml", ".ini": "ini",
    ".xml": "xml", ".md": "markdown", ".mdx": "markdown",
    ".vue": "vue", ".svelte": "svelte", ".swift": "swift", ".dart": "dart",
    ".lua": "lua", ".r": "r", ".scala": "scala", ".clj": "clojure",
    ".ex": "elixir", ".exs": "elixir", ".erl": "erlang", ".hs": "haskell",
    ".tf": "hcl", ".proto": "protobuf", ".graphql": "graphql", ".gql": "graphql",
}


class FileError(ValueError):
    """User-facing file handling error."""


def detect_language(ext: str) -> str:
    return LANGUAGE_BY_EXT.get((ext or "").lower(), "")


def detect_file_type(filename: str, mime: str | None = None) -> dict[str, Any]:
    """Classify a file into a category used by the context builder and the UI."""
    name = os.path.basename(filename or "")
    ext = os.path.splitext(name)[1].lower()
    mime = (mime or "").lower()

    if not ext and name.lower() in {"dockerfile", "makefile", "procfile", "rakefile", "gemfile"}:
        return {"category": "code", "ext": ext, "language": "bash"}

    if ext in ARCHIVE_EXTENSIONS or "zip" in mime:
        return {"category": "archive", "ext": ext, "language": ""}
    if ext in IMAGE_EXTENSIONS or mime.startswith("image/"):
        return {"category": "image", "ext": ext, "language": ""}
    if ext in DOCUMENT_EXTENSIONS or "pdf" in mime or "officedocument" in mime:
        return {"category": "document", "ext": ext, "language": ""}
    if ext in CODE_EXTENSIONS:
        return {"category": "code", "ext": ext, "language": detect_language(ext)}
    if ext in TEXT_EXTENSIONS or mime.startswith("text/"):
        return {"category": "text", "ext": ext, "language": detect_language(ext)}
    return {"category": "binary" if ext in SKIP_EXTENSIONS else "text", "ext": ext, "language": ""}


def is_sensitive_file(path: str) -> bool:
    """True when a path looks like it holds secrets."""
    norm = (path or "").replace("\\", "/").lower()
    base = norm.rsplit("/", 1)[-1]
    ext = os.path.splitext(base)[1]
    if ext in SENSITIVE_EXTENSIONS:
        return True
    if base == ".env" or base.startswith(".env."):
        return True
    if base.endswith(".env"):
        return True
    for pattern in SENSITIVE_NAME_PATTERNS:
        if pattern in base:
            return True
    if "/.ssh/" in norm or "/.aws/" in norm or "/.gnupg/" in norm:
        return True
    return False


def is_ignorable_path(path: str) -> bool:
    """True for build output, vendored dependencies and binaries."""
    norm = (path or "").replace("\\", "/").lower()
    parts = [p for p in norm.split("/") if p]
    for part in parts[:-1]:
        if part in SKIP_DIR_PARTS:
            return True
    base = parts[-1] if parts else norm
    ext = os.path.splitext(base)[1]
    if ext in SKIP_EXTENSIONS:
        return True
    if base.endswith(".min.js") or base.endswith(".min.css"):
        return True
    return False


def sanitize_filename(filename: str) -> str:
    base = os.path.basename((filename or "file").replace("\\", "/"))
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base).strip("._") or "file"
    return base[:120]


def guess_mime(ext: str) -> str:
    return _MIME_BY_EXT.get((ext or "").lower(), "application/octet-stream")


def _read_upload(file_storage: BinaryIO) -> bytes:
    """Read an upload exactly once, tolerating already-consumed streams."""
    try:
        file_storage.seek(0)
    except Exception:
        pass
    data = file_storage.read()
    if isinstance(data, str):
        data = data.encode("utf-8", errors="replace")
    return data or b""


def save_uploaded_file(
    file_storage,
    user_id: str,
    conversation_id: str | None = None,
    project_id: str | None = None,
    allow_sensitive: bool = False,
) -> dict[str, Any]:
    """Persist one uploaded file and register it against the owning user."""
    original_name = getattr(file_storage, "filename", "") or "file"
    content = _read_upload(file_storage)
    size = len(content)

    if size == 0:
        raise FileError(f"'{original_name}' is empty")
    if size > MAX_FILE_SIZE:
        raise FileError(
            f"'{original_name}' is {size // (1024 * 1024)}MB; the limit is {MAX_FILE_SIZE // (1024 * 1024)}MB"
        )
    if is_sensitive_file(original_name) and not allow_sensitive:
        raise FileError(
            f"'{original_name}' looks like a secrets file and was not uploaded. "
            "Rename it or re-upload with 'allow sensitive files' enabled if this is intentional."
        )

    file_id = f"file_{uuid.uuid4().hex[:12]}"
    declared_mime = getattr(file_storage, "content_type", None) or None
    info = detect_file_type(original_name, declared_mime)
    ext = info["ext"]

    if project_id:
        rel_dir = os.path.join("projects", project_id)
    elif conversation_id:
        rel_dir = os.path.join("conversations", conversation_id)
    else:
        rel_dir = os.path.join("users", user_id or "shared")

    target_dir = os.path.join(UPLOAD_DIR, rel_dir)
    os.makedirs(target_dir, exist_ok=True)
    disk_path = os.path.join(target_dir, f"{file_id}_{sanitize_filename(original_name)}")

    with open(disk_path, "wb") as fh:
        fh.write(content)

    mime = declared_mime or guess_mime(ext)
    now = datetime.now(timezone.utc).timestamp()

    db_execute(
        "INSERT INTO files (id, user_id, conversation_id, filename, mime, size, uploaded_at, path,"
        " file_type, processed) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (file_id, user_id, conversation_id, original_name, mime, size, now, disk_path,
         info["category"], 0),
    )

    return {
        "id": file_id,
        "filename": original_name,
        "mime": mime,
        "size": size,
        "file_type": info["category"],
        "language": info["language"],
        "ext": ext,
        "path": disk_path,
        "user_id": user_id,
        "conversation_id": conversation_id,
        "project_id": project_id,
    }


def save_generated_file(
    filename: str,
    content: str | bytes,
    user_id: str,
    conversation_id: str | None = None,
    project_id: str | None = None,
) -> dict[str, Any]:
    """Persist AI-generated file content and register it against the owning user.

    Content can be text (str) or binary (bytes). The filename must include
    the correct extension - detection is based on the extension.
    """
    if not filename or not filename.strip():
        raise FileError("Filename is required")

    original_name = filename.strip()
    ext = os.path.splitext(original_name)[1].lower()

    if isinstance(content, str):
        data = content.encode("utf-8", errors="replace")
    else:
        data = content or b""

    size = len(data)
    if size == 0:
        raise FileError(f"'{original_name}' is empty")

    info = detect_file_type(original_name, None)

    if project_id:
        rel_dir = os.path.join("projects", project_id)
    elif conversation_id:
        rel_dir = os.path.join("conversations", conversation_id)
    else:
        rel_dir = os.path.join("users", user_id or "shared", "generated")

    target_dir = os.path.join(UPLOAD_DIR, rel_dir)
    os.makedirs(target_dir, exist_ok=True)

    file_id = f"gen_{uuid.uuid4().hex[:12]}"
    disk_path = os.path.join(target_dir, f"{file_id}_{sanitize_filename(original_name)}")

    with open(disk_path, "wb") as fh:
        fh.write(data)

    mime = guess_mime(ext)
    now = datetime.now(timezone.utc).timestamp()

    db_execute(
        "INSERT INTO files (id, user_id, conversation_id, filename, mime, size, uploaded_at, path,"
        " file_type, processed) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (file_id, user_id, conversation_id, original_name, mime, size, now, disk_path,
         info["category"], 1),
    )

    return {
        "id": file_id,
        "filename": original_name,
        "mime": mime,
        "size": size,
        "file_type": info["category"],
        "language": info["language"],
        "ext": ext,
        "path": disk_path,
        "user_id": user_id,
        "conversation_id": conversation_id,
        "project_id": project_id,
    }


def _extract_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[page {index + 1}]\n{text}")
    return "\n\n".join(pages)


def _extract_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    blocks = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def _extract_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    blocks = []
    for index, slide in enumerate(prs.slides):
        parts = [f"[slide {index + 1}]"]
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                parts.append(text)
        if len(parts) > 1:
            blocks.append("\n".join(parts))
    return "\n\n".join(blocks)


def _extract_html(path: str) -> str:
    from bs4 import BeautifulSoup

    with open(path, encoding="utf-8", errors="replace") as fh:
        soup = BeautifulSoup(fh.read(), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text(separator="\n", strip=True)


def _extract_ipynb(path: str) -> str:
    import json as _json

    with open(path, encoding="utf-8", errors="replace") as fh:
        try:
            nb = _json.load(fh)
        except Exception:
            fh.seek(0)
            return fh.read()
    blocks = []
    for cell in nb.get("cells", []):
        source = cell.get("source", [])
        text = "".join(source) if isinstance(source, list) else str(source)
        if not text.strip():
            continue
        if cell.get("cell_type") == "code":
            blocks.append("```python\n" + text + "\n```")
        else:
            blocks.append(text)
    return "\n\n".join(blocks)


def extract_text_from_file(path: str, limit: int = MAX_TEXT_EXTRACT_CHARS) -> str:
    """Best-effort plain-text extraction. Never raises."""
    if not path or not os.path.isfile(path):
        return ""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            text = _extract_pdf(path)
        elif ext in {".docx", ".doc"}:
            text = _extract_docx(path)
        elif ext in {".pptx", ".ppt"}:
            text = _extract_pptx(path)
        elif ext in {".html", ".htm"}:
            text = _extract_html(path)
        elif ext == ".ipynb":
            text = _extract_ipynb(path)
        elif ext in IMAGE_EXTENSIONS:
            return ""
        else:
            with open(path, encoding="utf-8", errors="replace") as fh:
                text = fh.read(limit + 1)
    except Exception as exc:  # noqa: BLE001 - extraction must not break a request
        log.warning("Text extraction failed for %s: %s", path, exc)
        return f"[Could not read this file: {exc}]"

    text = text.replace("\x00", "")
    if len(text) > limit:
        text = text[:limit] + f"\n\n[... truncated at {limit} characters ...]"
    return text


def _safe_zip_target(base_dir: str, entry_name: str) -> str | None:
    """Resolve a ZIP entry to a path inside base_dir, or None if it escapes."""
    cleaned = (entry_name or "").replace("\\", "/").lstrip("/")
    if not cleaned or cleaned.endswith("/"):
        return None
    if any(part == ".." for part in cleaned.split("/")):
        return None
    if re.match(r"^[A-Za-z]:", cleaned):
        return None
    target = os.path.normpath(os.path.join(base_dir, cleaned))
    base_real = os.path.realpath(base_dir)
    target_real = os.path.realpath(os.path.dirname(target))
    if not (target_real == base_real or target_real.startswith(base_real + os.sep)):
        return None
    return target


def process_zip(
    zip_path: str,
    project_id: str,
    user_id: str,
    conversation_id: str | None = None,
    allow_sensitive: bool = False,
) -> dict[str, Any]:
    """Extract a ZIP into a project sandbox and register its source files.

    Blocks path traversal, absolute paths, secrets, vendored directories and
    zip bombs. Returns a report describing what was kept and what was skipped.
    """
    if not project_id:
        raise FileError("A project is required before extracting an archive")

    base_extract = os.path.join(UPLOAD_DIR, "projects", project_id, "extracted")
    os.makedirs(base_extract, exist_ok=True)

    kept: list[dict[str, Any]] = []
    skipped_sensitive: list[str] = []
    skipped_ignored: list[str] = []
    skipped_large: list[str] = []
    total_uncompressed = 0

    try:
        with zipfile.ZipFile(zip_path, "r") as zf:
            entries = [i for i in zf.infolist() if not i.is_dir()]
            if len(entries) > MAX_ZIP_ENTRIES:
                raise FileError(
                    f"Archive holds {len(entries)} files; the limit is {MAX_ZIP_ENTRIES}"
                )

            rows_files: list[tuple] = []
            rows_project: list[tuple] = []
            now = datetime.now(timezone.utc).timestamp()

            for info in entries:
                name = info.filename
                if is_sensitive_file(name) and not allow_sensitive:
                    skipped_sensitive.append(name)
                    continue
                if is_ignorable_path(name):
                    skipped_ignored.append(name)
                    continue
                if info.file_size > MAX_ZIP_ENTRY_SIZE:
                    skipped_large.append(name)
                    continue
                total_uncompressed += info.file_size
                if total_uncompressed > MAX_ZIP_TOTAL_UNCOMPRESSED:
                    raise FileError("Archive expands beyond the 300MB safety limit")

                target = _safe_zip_target(base_extract, name)
                if target is None:
                    skipped_ignored.append(name)
                    continue

                os.makedirs(os.path.dirname(target), exist_ok=True)
                with zf.open(info, "r") as src, open(target, "wb") as dst:
                    dst.write(src.read(MAX_ZIP_ENTRY_SIZE + 1))

                rel_path = os.path.relpath(target, base_extract).replace("\\", "/")
                ext = os.path.splitext(rel_path)[1].lower()
                kind = detect_file_type(rel_path, None)
                file_id = f"file_{uuid.uuid4().hex[:12]}"

                content = ""
                if kind["category"] in {"code", "text", "document"}:
                    content = extract_text_from_file(target, limit=120_000)

                rows_files.append(
                    (file_id, user_id, conversation_id, rel_path, guess_mime(ext), info.file_size,
                     now, target, kind["category"], 1)
                )
                rows_project.append(
                    (f"pf_{uuid.uuid4().hex[:12]}", project_id, file_id, rel_path,
                     os.path.basename(rel_path), ext, kind["category"], content, info.file_size)
                )
                kept.append({
                    "id": file_id,
                    "path": rel_path,
                    "file_type": kind["category"],
                    "language": kind["language"],
                    "size": info.file_size,
                })

            if rows_files:
                from server.storage.db import executemany as db_executemany

                db_executemany(
                    "INSERT INTO files (id, user_id, conversation_id, filename, mime, size,"
                    " uploaded_at, path, file_type, processed)"
                    " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    rows_files,
                )
                db_executemany(
                    "INSERT INTO project_files (id, project_id, file_id, path, name, ext,"
                    " file_type, content, size) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    rows_project,
                )

    except zipfile.BadZipFile:
        raise FileError("That archive is not a valid ZIP file")

    return {
        "project_id": project_id,
        "kept": kept,
        "file_count": len(kept),
        "skipped_sensitive": skipped_sensitive,
        "skipped_ignored": skipped_ignored[:50],
        "skipped_ignored_count": len(skipped_ignored),
        "skipped_large": skipped_large,
    }


def get_file_record(file_id: str) -> dict[str, Any] | None:
    return db_query("SELECT * FROM files WHERE id = ?", (file_id,), one=True)
